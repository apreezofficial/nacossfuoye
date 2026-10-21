import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\USER\Downloads\DOC-20260909-WA0019.xlsx')
sheet = wb['Sheet1']

levels_data = []
current_level = None

for r in range(1, sheet.max_row + 1):
    c1 = sheet.cell(r, 1).value
    c2 = sheet.cell(r, 2).value
    c3 = sheet.cell(r, 3).value
    c4 = sheet.cell(r, 4).value
    
    if c1 and ('100' in str(c1) or '200' in str(c1) or '300' in str(c1) or '400' in str(c1)) and not c2:
        digits = ''.join(filter(str.isdigit, str(c1)))[:3]
        current_level = {
            'level': digits,
            'title': f'{digits} Level',
            'courses': []
        }
        levels_data.append(current_level)
        continue
    
    if c1 and 'Course Code' in str(c1):
        continue
        
    if current_level and c1 and c2:
        col3 = c3
        col4 = c4
        
        if isinstance(col3, (int, float)) or (isinstance(col3, str) and str(col3).strip().isdigit()):
            units = int(col3)
            status = str(col4).strip() if col4 is not None else ''
        elif isinstance(col4, (int, float)) or (isinstance(col4, str) and str(col4).strip().isdigit()):
            units = int(col4)
            status = str(col3).strip() if col3 is not None else ''
        else:
            units = 0
            status = ''
            
        code = str(c1).strip()
        title = str(c2).strip()
        current_level['courses'].append({
            'code': code,
            'title': title,
            'units': units,
            'status': status
        })

for lvl in levels_data:
    tot_units = sum(c['units'] for c in lvl['courses'])
    lvl['totalUnits'] = tot_units
    lvl['totalCourses'] = len(lvl['courses'])
    print(f"{lvl['title']}: {len(lvl['courses'])} courses, {tot_units} units")

with open('src/data/courses.json', 'w') as f:
    json.dump(levels_data, f, indent=2)

print('Saved to src/data/courses.json')
